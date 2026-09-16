import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
rows=json.load(open("rows.json")); needed=json.load(open("needed.json"))
SUFFIX={'Harness':'-Harness','Section':'-Section','Activity':'-Act','Data Transform':'-DT','RDB List':'-SQL','Flow':'-Flow'}
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
hf=PatternFill("solid",fgColor="000000"); hfont=Font(name="Arial",bold=True,color="FFFFFF",size=11)
wb=openpyxl.Workbook()
ws=wb.active; ws.title="Rule Structure"
rootf=PatternFill("solid",fgColor="D9D9D9"); sectf=PatternFill("solid",fgColor="FCE4D6")
dtf=PatternFill("solid",fgColor="E2EFDA"); dupf=PatternFill("solid",fgColor="F2F2F2")
for c,h in enumerate(["No","Nama Rule","Shape ID","Jenis Rule","XML"],1):
    x=ws.cell(1,c,h); x.fill=hf; x.font=hfont; x.border=bd; x.alignment=Alignment(horizontal="left",vertical="center")
r=2
for level,no,name,jenis,dup in rows:
    label=("   "*level)+name+(f"   \u2191 lihat No {dup}" if dup else "")
    xmlf="InboxAutoClaim-Harness.xml" if level==0 else f"{no} {name}{SUFFIX[jenis]}.xml"
    ws.cell(r,1,no); ws.cell(r,2,label); ws.cell(r,3,"Harness1" if level==0 else "")
    ws.cell(r,4,jenis); ws.cell(r,5,("" if dup else "\U0001F5CE ")+xmlf)
    bold=(not dup) and (level==0 or jenis in ("Section","Harness"))
    if dup: fill=dupf
    elif level==0: fill=rootf
    elif jenis=="Section": fill=sectf
    elif jenis=="Data Transform": fill=dtf
    else: fill=None
    for c in range(1,6):
        cc=ws.cell(r,c); cc.border=bd; cc.alignment=Alignment(vertical="center")
        cc.font=Font(name="Arial",bold=bold,italic=bool(dup),size=11,color=("808080" if dup else "000000"))
        if fill: cc.fill=fill
    r+=1
for col,w in zip("ABCDE",[20,64,10,16,66]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"; ws.row_dimensions[1].height=22

ws2=wb.create_sheet("XML needed")
actf=PatternFill("solid",fgColor="FFF2CC")
for c,h in enumerate(["No","Nama Rule","Jenis Rule","XML File Needed","Direferensikan Oleh","Bisa Diurai Lagi?"],1):
    x=ws2.cell(1,c,h); x.fill=hf; x.font=hfont; x.border=bd; x.alignment=Alignment(horizontal="left",vertical="center")
r=2
for i,(t,n,exp,par) in enumerate(needed,1):
    expand="Ya (punya turunan)" if t in ("Activity","Section","Flow","Data Transform") else "Tidak (leaf/query)"
    ws2.cell(r,1,i); ws2.cell(r,2,n); ws2.cell(r,3,t); ws2.cell(r,4,exp); ws2.cell(r,5,", ".join(par)); ws2.cell(r,6,expand)
    for c in range(1,7):
        cc=ws2.cell(r,c); cc.border=bd; cc.font=Font(name="Arial",size=11); cc.alignment=Alignment(vertical="center")
        if t in ("Activity","Data Transform","Section","Flow"): cc.fill=actf
    r+=1
for col,w in zip("ABCDEF",[6,44,14,50,50,20]): ws2.column_dimensions[col].width=w
ws2.freeze_panes="A2"; ws2.row_dimensions[1].height=22
wb.save("/mnt/user-data/outputs/InboxAutoClaim-Rule-Structure.xlsx")
print("saved. tree",len(rows),"needed",len(needed))
